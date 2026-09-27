from pathlib import Path

from pypdf import PdfReader
from docx import Document


SUPPORTED_FORMATS = {".pdf", ".docx", ".txt"}


def extract_text(uploaded_file):
    # Identify the uploaded file format from its extension.
    extension = Path(uploaded_file.name).suffix.lower()

    if extension not in SUPPORTED_FORMATS:
        raise ValueError(
            "Only PDF, DOCX and TXT files are supported."
        )

    if extension == ".pdf":
        # Extract text from every page of a PDF.
        pages = [
            page.extract_text()
            for page in PdfReader(uploaded_file).pages
        ]
        text = "\n".join(page for page in pages if page)

    elif extension == ".docx":
        # Extract paragraphs and table content from a DOCX file.
        document = Document(uploaded_file)

        parts = [
            paragraph.text.strip()
            for paragraph in document.paragraphs
            if paragraph.text.strip()
        ]

        for table in document.tables:
            for row in table.rows:
                values = [
                    cell.text.strip()
                    for cell in row.cells
                    if cell.text.strip()
                ]

                if values:
                    parts.append(" | ".join(values))

        text = "\n".join(parts)

    else:
        # Decode plain-text files as UTF-8.
        text = uploaded_file.getvalue().decode(
            "utf-8",
            errors="ignore"
        )

    # Remove null characters and unnecessary whitespace.
    text = text.replace("\x00", " ").strip()

    if len(text) < 40:
        raise ValueError(
            "Very little text could be extracted. "
            "Scanned PDFs require OCR."
        )

    return text
